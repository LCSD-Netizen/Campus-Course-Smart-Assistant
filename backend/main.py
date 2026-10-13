from __future__ import annotations

import asyncio
import hashlib
import hmac
import json
import os
import re
import secrets
import uuid
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
from typing import Any, AsyncIterator

import httpx
import jwt
from dotenv import load_dotenv
from fastapi import Depends, FastAPI, File, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse
from fastapi.security import HTTPAuthorizationCredentials, HTTPBearer
from pydantic import BaseModel, Field
from sqlalchemy import Boolean, Date, DateTime, ForeignKey, Integer, String, Text, create_engine, delete, func, select
from sqlalchemy.orm import DeclarativeBase, Mapped, Session, mapped_column, relationship, sessionmaker

load_dotenv()
ROOT = Path(__file__).resolve().parents[1]
DATA_DIR = ROOT / "data"
DATA_DIR.mkdir(exist_ok=True)
DATABASE_URL = os.getenv("DATABASE_URL", "sqlite:///./data/campus.db")
if DATABASE_URL.startswith("sqlite"):
    engine = create_engine(DATABASE_URL, connect_args={"check_same_thread": False})
else:
    engine = create_engine(DATABASE_URL, pool_pre_ping=True)
SessionLocal = sessionmaker(bind=engine, expire_on_commit=False)
SECRET_KEY = os.getenv("SECRET_KEY", "development-only-change-this-secret")
ALGORITHM = "HS256"
bearer = HTTPBearer(auto_error=False)


class Base(DeclarativeBase):
    pass


class User(Base):
    __tablename__ = "users"
    id: Mapped[int] = mapped_column(Integer, primary_key=True)
    username: Mapped[str] = mapped_column(String(50), unique=True, index=True)
    display_name: Mapped[str] = mapped_column(String(80), default="同学")
    password_hash: Mapped[str] = mapped_column(String(256))
    role: Mapped[str] = mapped_column(String(20), default="student")
    department: Mapped[str] = mapped_column(String(100), default="计算机科学与技术")
    created_at: Mapped[datetime] = mapped_column(DateTime, default=datetime.utcnow)


class Course(Base):
    __tablename__ = "courses"
    id: Mapped[int] = mapped_column(Integer, primary_key=True)
    name: Mapped[str] = mapped_column(String(120), index=True)
    teacher: Mapped[str] = mapped_column(String(80), default="任课教师")
    room: Mapped[str] = mapped_column(String(80), default="待安排")
    weekday: Mapped[int] = mapped_column(Integer, default=1)
    start_time: Mapped[str] = mapped_column(String(8), default="08:00")
    end_time: Mapped[str] = mapped_column(String(8), default="09:40")
    weeks: Mapped[str] = mapped_column(String(80), default="1-16周")
    credits: Mapped[int] = mapped_column(Integer, default=3)
    color: Mapped[str] = mapped_column(String(20), default="blue")


class Enrollment(Base):
    __tablename__ = "enrollments"
    id: Mapped[int] = mapped_column(Integer, primary_key=True)
    user_id: Mapped[int] = mapped_column(ForeignKey("users.id"), index=True)
    course_id: Mapped[int] = mapped_column(ForeignKey("courses.id"), index=True)


class Task(Base):
    __tablename__ = "tasks"
    id: Mapped[int] = mapped_column(Integer, primary_key=True)
    user_id: Mapped[int] = mapped_column(ForeignKey("users.id"), index=True)
    course_id: Mapped[int | None] = mapped_column(ForeignKey("courses.id"), nullable=True)
    title: Mapped[str] = mapped_column(String(200))
    kind: Mapped[str] = mapped_column(String(40), default="作业")
    due_date: Mapped[date] = mapped_column(Date)
    completed: Mapped[bool] = mapped_column(Boolean, default=False)


class Plan(Base):
    __tablename__ = "study_plans"
    id: Mapped[int] = mapped_column(Integer, primary_key=True)
    user_id: Mapped[int] = mapped_column(ForeignKey("users.id"), index=True)
    title: Mapped[str] = mapped_column(String(180))
    content: Mapped[str] = mapped_column(Text)
    created_at: Mapped[datetime] = mapped_column(DateTime, default=datetime.utcnow)
    due_date: Mapped[date | None] = mapped_column(Date, nullable=True)


class Document(Base):
    __tablename__ = "documents"
    id: Mapped[int] = mapped_column(Integer, primary_key=True)
    course_id: Mapped[int | None] = mapped_column(ForeignKey("courses.id"), nullable=True)
    filename: Mapped[str] = mapped_column(String(255))
    content: Mapped[str] = mapped_column(Text)
    uploaded_by: Mapped[int] = mapped_column(ForeignKey("users.id"))
    created_at: Mapped[datetime] = mapped_column(DateTime, default=datetime.utcnow)


class Credentials(BaseModel):
    username: str
    password: str


class RegisterBody(Credentials):
    display_name: str = "同学"


class ChatBody(BaseModel):
    message: str = Field(min_length=1, max_length=4000)


class PlanBody(BaseModel):
    title: str
    content: str
    due_date: date | None = None


class TaskBody(BaseModel):
    title: str
    kind: str = "作业"
    due_date: date
    course_id: int | None = None


class UserRoleBody(BaseModel):
    role: str


class CourseAccessBody(BaseModel):
    course_ids: list[int]


def password_hash(password: str, salt: str | None = None) -> str:
    salt = salt or secrets.token_hex(16)
    digest = hashlib.pbkdf2_hmac("sha256", password.encode(), bytes.fromhex(salt), 260_000).hex()
    return f"{salt}${digest}"


def verify_password(password: str, stored: str) -> bool:
    salt, digest = stored.split("$", 1)
    return hmac.compare_digest(password_hash(password, salt).split("$", 1)[1], digest)


def token_for(user: User) -> str:
    return jwt.encode({"sub": str(user.id), "role": user.role, "exp": datetime.now(timezone.utc) + timedelta(hours=24)}, SECRET_KEY, algorithm=ALGORITHM)


def get_db():
    with SessionLocal() as db:
        yield db


def get_user(credentials: HTTPAuthorizationCredentials | None = Depends(bearer), db: Session = Depends(get_db)) -> User:
    if not credentials:
        raise HTTPException(401, "请先登录")
    try:
        payload = jwt.decode(credentials.credentials, SECRET_KEY, algorithms=[ALGORITHM])
        user = db.get(User, int(payload["sub"]))
    except Exception as exc:
        raise HTTPException(401, "登录已过期，请重新登录") from exc
    if not user:
        raise HTTPException(401, "用户不存在")
    return user


def admin_only(user: User = Depends(get_user)) -> User:
    if user.role != "admin":
        raise HTTPException(403, "需要管理员权限")
    return user


def course_json(c: Course) -> dict[str, Any]:
    return {"id": c.id, "name": c.name, "teacher": c.teacher, "room": c.room, "weekday": c.weekday, "start_time": c.start_time, "end_time": c.end_time, "weeks": c.weeks, "credits": c.credits, "color": c.color}


def task_json(t: Task, db: Session) -> dict[str, Any]:
    c = db.get(Course, t.course_id) if t.course_id else None
    return {"id": t.id, "title": t.title, "kind": t.kind, "due_date": t.due_date.isoformat(), "completed": t.completed, "course_id": t.course_id, "course_name": c.name if c else "个人计划", "days_left": (t.due_date - date.today()).days}


def signed_user_context(user_id: int) -> str:
    raw = str(user_id)
    signature = hmac.new(SECRET_KEY.encode(), raw.encode(), hashlib.sha256).hexdigest()
    return f"{raw}.{signature}"


async def call_mcp_tool(name: str, arguments: dict[str, Any]) -> str | None:
    url = os.getenv("MCP_URL", "http://127.0.0.1:8001/mcp")
    try:
        from mcp import ClientSession
        from mcp.client.streamable_http import streamablehttp_client
        async with streamablehttp_client(url) as (read_stream, write_stream, _):
            async with ClientSession(read_stream, write_stream) as session:
                await session.initialize()
                result = await session.call_tool(name, arguments)
                return "\n".join(item.text for item in result.content if hasattr(item, "text"))
    except Exception:
        return None


def seed_data():
    Base.metadata.create_all(engine)
    with SessionLocal() as db:
        if not db.scalar(select(User).where(User.username == "demo")):
            student = User(username="demo", display_name="林知夏", password_hash=password_hash("demo1234"), role="student")
            admin = User(username="admin", display_name="管理员", password_hash=password_hash("admin1234"), role="admin", department="校园管理")
            db.add_all([student, admin]); db.flush()
            sample = [
                Course(name="数据结构与算法", teacher="王建国", room="理科楼 302", weekday=1, start_time="08:00", end_time="09:40", color="blue"),
                Course(name="计算机网络", teacher="李文清", room="信息楼 205", weekday=2, start_time="10:00", end_time="11:40", color="green"),
                Course(name="操作系统原理", teacher="陈思远", room="理科楼 410", weekday=3, start_time="14:00", end_time="15:40", credits=4, color="orange"),
                Course(name="数据库系统", teacher="赵明", room="信息楼 318", weekday=4, start_time="08:00", end_time="09:40", color="purple"),
                Course(name="软件工程", teacher="刘老师", room="科创中心 201", weekday=5, start_time="10:00", end_time="11:40", credits=2, color="green"),
                Course(name="大学英语（四）", teacher="周敏", room="外语楼 106", weekday=5, start_time="14:00", end_time="15:40", credits=2, color="blue"),
            ]
            db.add_all(sample); db.flush()
            for course in sample: db.add(Enrollment(user_id=student.id, course_id=course.id))
            db.add_all([
                Task(user_id=student.id, course_id=sample[0].id, title="第 5 章编程作业", kind="作业", due_date=date.today()+timedelta(days=1)),
                Task(user_id=student.id, course_id=sample[1].id, title="Wireshark 实验报告", kind="实验报告", due_date=date.today()+timedelta(days=3)),
                Task(user_id=student.id, course_id=sample[2].id, title="期中测验", kind="考试", due_date=date.today()+timedelta(days=5)),
                Task(user_id=student.id, course_id=sample[3].id, title="SQL 查询练习", kind="作业", due_date=date.today()+timedelta(days=7)),
            ])
            db.add_all([
                Document(course_id=sample[0].id, filename="数据结构课程讲义.pdf", uploaded_by=admin.id, content="第三章 栈与队列。栈是一种后进先出（LIFO）的线性数据结构，基本操作包括入栈、出栈和取栈顶。队列是一种先进先出（FIFO）的线性数据结构，循环队列通过取模运算复用数组空间。栈的典型应用包括括号匹配、表达式求值和函数调用；队列常用于任务调度和广度优先搜索。"),
                Document(course_id=sample[1].id, filename="计算机网络课程讲义.pdf", uploaded_by=admin.id, content="传输层提供端到端通信。TCP 面向连接，通过序列号、确认应答和重传实现可靠、有序的数据传输；UDP 无连接，不保证送达和顺序，但开销较小、延迟较低。文件传输通常使用 TCP，实时音视频和 DNS 常使用 UDP。"),
                Document(course_id=sample[2].id, filename="操作系统复习提纲.pdf", uploaded_by=admin.id, content="操作系统期中复习重点：进程与线程的状态转换和调度算法；同步互斥、信号量与经典同步问题；内存分页、虚拟内存和页面置换算法。"),
            ])
            db.commit()


seed_data()
app = FastAPI(title="知课校园课程智能助手 API", version="1.0.0")
origins = os.getenv("FRONTEND_ORIGINS", "http://localhost:5173,http://127.0.0.1:5173").split(",")
app.add_middleware(CORSMiddleware, allow_origins=origins, allow_credentials=True, allow_methods=["*"], allow_headers=["*"])


@app.get("/api/health")
def health():
    return {"status": "ok", "agent_configured": bool(os.getenv("OPENAI_API_KEY"))}


@app.post("/api/auth/register")
def register(body: RegisterBody, db: Session = Depends(get_db)):
    if len(body.password) < 8: raise HTTPException(400, "密码至少需要 8 位")
    if db.scalar(select(User).where(User.username == body.username)): raise HTTPException(409, "用户名已存在")
    user = User(username=body.username, display_name=body.display_name, password_hash=password_hash(body.password))
    db.add(user); db.commit(); db.refresh(user)
    return {"access_token": token_for(user), "token_type": "bearer", "user": {"id": user.id, "username": user.username, "display_name": user.display_name, "role": user.role}}


@app.post("/api/auth/login")
def login(body: Credentials, db: Session = Depends(get_db)):
    user = db.scalar(select(User).where(User.username == body.username))
    if not user or not verify_password(body.password, user.password_hash): raise HTTPException(401, "用户名或密码错误")
    return {"access_token": token_for(user), "token_type": "bearer", "user": {"id": user.id, "username": user.username, "display_name": user.display_name, "role": user.role}}


@app.get("/api/auth/me")
def me(user: User = Depends(get_user)):
    return {"id": user.id, "username": user.username, "display_name": user.display_name, "role": user.role, "department": user.department}


@app.get("/api/dashboard")
def dashboard(user: User = Depends(get_user), db: Session = Depends(get_db)):
    courses = db.scalars(select(Course).join(Enrollment).where(Enrollment.user_id == user.id)).all()
    tasks = db.scalars(select(Task).where(Task.user_id == user.id).order_by(Task.due_date)).all()
    plans = db.scalars(select(Plan).where(Plan.user_id == user.id).order_by(Plan.created_at.desc()).limit(5)).all()
    enrolled = [c.id for c in courses]
    docs = db.scalar(select(func.count(Document.id)).where((Document.course_id.is_(None)) | (Document.course_id.in_(enrolled if enrolled else [-1])))) or 0
    return {"courses": [course_json(c) for c in courses], "tasks": [task_json(t, db) for t in tasks], "plans": [{"id": p.id, "title": p.title, "content": p.content, "due_date": p.due_date.isoformat() if p.due_date else None, "created_at": p.created_at.isoformat()} for p in plans], "stats": {"courses": len(courses), "pending_tasks": sum(not t.completed for t in tasks), "documents": docs}}


@app.get("/api/schedule")
def schedule(user: User = Depends(get_user), db: Session = Depends(get_db)):
    courses = db.scalars(select(Course).join(Enrollment).where(Enrollment.user_id == user.id).order_by(Course.weekday, Course.start_time)).all()
    return [course_json(c) for c in courses]


@app.get("/api/courses")
def list_courses(user: User = Depends(get_user), db: Session = Depends(get_db)):
    query = select(Course).order_by(Course.weekday, Course.start_time)
    if user.role != "admin": query = query.join(Enrollment).where(Enrollment.user_id == user.id)
    return [course_json(c) for c in db.scalars(query).all()]


@app.get("/api/tasks")
def list_tasks(user: User = Depends(get_user), db: Session = Depends(get_db)):
    tasks = db.scalars(select(Task).where(Task.user_id == user.id).order_by(Task.due_date)).all()
    return [task_json(t, db) for t in tasks]


@app.post("/api/tasks")
def add_task(body: TaskBody, user: User = Depends(get_user), db: Session = Depends(get_db)):
    if body.course_id and not db.scalar(select(Enrollment).where(Enrollment.user_id == user.id, Enrollment.course_id == body.course_id)):
        raise HTTPException(403, "无权为该课程创建任务")
    row = Task(user_id=user.id, course_id=body.course_id, title=body.title, kind=body.kind, due_date=body.due_date)
    db.add(row); db.commit(); db.refresh(row)
    return task_json(row, db)


@app.patch("/api/tasks/{task_id}/complete")
def complete_task(task_id: int, user: User = Depends(get_user), db: Session = Depends(get_db)):
    row = db.scalar(select(Task).where(Task.id == task_id, Task.user_id == user.id))
    if not row: raise HTTPException(404, "任务不存在")
    row.completed = not row.completed; db.commit()
    return task_json(row, db)


@app.get("/api/plans")
def list_plans(user: User = Depends(get_user), db: Session = Depends(get_db)):
    rows = db.scalars(select(Plan).where(Plan.user_id == user.id).order_by(Plan.created_at.desc())).all()
    return [{"id": p.id, "title": p.title, "content": p.content, "due_date": p.due_date.isoformat() if p.due_date else None, "created_at": p.created_at.isoformat()} for p in rows]


@app.post("/api/plans")
def create_plan(body: PlanBody, user: User = Depends(get_user), db: Session = Depends(get_db)):
    row = Plan(user_id=user.id, title=body.title, content=body.content, due_date=body.due_date)
    db.add(row); db.commit(); db.refresh(row)
    return {"id": row.id, "title": row.title, "content": row.content, "due_date": row.due_date.isoformat() if row.due_date else None}


def split_chunks(text: str, size: int = 450, overlap: int = 80) -> list[str]:
    text = re.sub(r"\s+", " ", text).strip()
    return [text[i:i + size] for i in range(0, len(text), size - overlap) if text[i:i + size].strip()]


async def embeddings(texts: list[str]) -> list[list[float]] | None:
    key = os.getenv("EMBEDDING_API_KEY")
    if not key: return None
    url = os.getenv("EMBEDDING_BASE_URL", "https://api.siliconflow.cn/v1").rstrip("/") + "/embeddings"
    try:
        async with httpx.AsyncClient(timeout=45) as client:
            response = await client.post(url, headers={"Authorization": f"Bearer {key}"}, json={"model": os.getenv("EMBEDDING_MODEL", "BAAI/bge-m3"), "input": texts})
            response.raise_for_status()
            return [item["embedding"] for item in response.json()["data"]]
    except Exception:
        return None


async def vector_add(doc: Document):
    chunks = split_chunks(doc.content)
    vectors = await embeddings(chunks)
    if vectors is None: return
    try:
        import chromadb
        path = Path(os.getenv("CHROMA_DIR", str(DATA_DIR / "chroma"))); path.mkdir(parents=True, exist_ok=True)
        client = chromadb.PersistentClient(path=str(path))
        collection = client.get_or_create_collection("campus_course_materials")
        collection.add(ids=[f"doc-{doc.id}-{i}" for i in range(len(chunks))], documents=chunks, embeddings=vectors, metadatas=[{"document_id": doc.id, "course_id": doc.course_id or 0, "filename": doc.filename} for _ in chunks])
    except Exception:
        return


async def search_materials(query: str, user: User, db: Session, limit: int = 4) -> list[dict[str, Any]]:
    enrolled = db.scalars(select(Enrollment.course_id).where(Enrollment.user_id == user.id)).all()
    docs = db.scalars(select(Document).where((Document.course_id.is_(None)) | (Document.course_id.in_(enrolled if enrolled else [-1])))).all()
    query_vector = await embeddings([query])
    if query_vector:
        try:
            import chromadb
            client = chromadb.PersistentClient(path=os.getenv("CHROMA_DIR", str(DATA_DIR / "chroma")))
            collection = client.get_or_create_collection("campus_course_materials")
            result = collection.query(query_embeddings=query_vector, n_results=limit, where={"course_id": {"$in": enrolled}} if enrolled else {"course_id": -1})
            doc_map = {d.id: d for d in docs}
            ids = result.get("ids", [[]])[0]
            chunks = result.get("documents", [[]])[0]
            rows = []
            for chunk_id, chunk in zip(ids, chunks):
                match = re.search(r"doc-(\d+)-", chunk_id)
                document_id = int(match.group(1)) if match else None
                if document_id in doc_map:
                    rows.append({"filename": doc_map[document_id].filename, "content": chunk})
            if rows: return rows[:limit]
        except Exception:
            pass
    terms = [t.lower() for t in re.findall(r"[\w\u4e00-\u9fff]+", query) if len(t) > 1]
    ranked = sorted(docs, key=lambda d: sum(d.content.lower().count(t) + d.filename.lower().count(t) for t in terms), reverse=True)
    return [{"filename": d.filename, "content": d.content[:650]} for d in ranked[:limit]]


@app.get("/api/documents")
def list_documents(user: User = Depends(get_user), db: Session = Depends(get_db)):
    if user.role == "admin":
        rows = db.scalars(select(Document).order_by(Document.created_at.desc())).all()
        return [{"id": d.id, "filename": d.filename, "course_id": d.course_id, "created_at": d.created_at.isoformat(), "characters": len(d.content)} for d in rows]
    enrolled = db.scalars(select(Enrollment.course_id).where(Enrollment.user_id == user.id)).all()
    rows = db.scalars(select(Document).where((Document.course_id.is_(None)) | (Document.course_id.in_(enrolled if enrolled else [-1]))).order_by(Document.created_at.desc())).all()
    return [{"id": d.id, "filename": d.filename, "course_id": d.course_id, "created_at": d.created_at.isoformat(), "characters": len(d.content)} for d in rows]


async def extract_text(filename: str, data: bytes) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix == ".pdf":
        from pypdf import PdfReader
        import io
        return "\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(data)).pages)
    if suffix in (".docx",):
        from docx import Document as WordDocument
        import io
        return "\n".join(p.text for p in WordDocument(io.BytesIO(data)).paragraphs)
    if suffix in (".pptx",):
        from pptx import Presentation
        import io
        prs = Presentation(io.BytesIO(data))
        return "\n".join(shape.text for slide in prs.slides for shape in slide.shapes if getattr(shape, "has_text_frame", False))
    if suffix == ".txt": return data.decode("utf-8", errors="ignore")
    raise HTTPException(400, "支持 PDF、DOCX、PPTX、TXT 格式")


@app.post("/api/admin/documents")
async def upload_document(file: UploadFile = File(...), course_id: int | None = None, user: User = Depends(admin_only), db: Session = Depends(get_db)):
    if Path(file.filename or "").suffix.lower() not in {".pdf", ".docx", ".pptx", ".txt"}: raise HTTPException(400, "不支持的文件格式")
    raw = await file.read()
    if len(raw) > 20 * 1024 * 1024: raise HTTPException(413, "单个文件不能超过 20MB")
    if course_id and not db.get(Course, course_id): raise HTTPException(404, "课程不存在")
    text = await extract_text(file.filename or "资料", raw)
    if not text.strip(): raise HTTPException(400, "文件中未能提取到文本内容")
    row = Document(filename=Path(file.filename or "资料").name, content=text[:2_000_000], course_id=course_id, uploaded_by=user.id)
    db.add(row); db.commit(); db.refresh(row)
    await vector_add(row)
    return {"id": row.id, "filename": row.filename, "characters": len(row.content), "chunks": len(split_chunks(row.content))}


@app.get("/api/admin/users")
def admin_users(user: User = Depends(admin_only), db: Session = Depends(get_db)):
    rows = db.scalars(select(User).order_by(User.id)).all()
    enrollments = db.scalars(select(Enrollment)).all()
    course_map: dict[int, list[int]] = {}
    for enrollment in enrollments: course_map.setdefault(enrollment.user_id, []).append(enrollment.course_id)
    return [{"id": u.id, "username": u.username, "display_name": u.display_name, "role": u.role, "department": u.department, "course_ids": course_map.get(u.id, []), "created_at": u.created_at.isoformat()} for u in rows]


@app.patch("/api/admin/users/{user_id}/role")
def set_user_role(user_id: int, body: UserRoleBody, user: User = Depends(admin_only), db: Session = Depends(get_db)):
    if body.role not in {"student", "admin"}: raise HTTPException(400, "无效角色")
    target = db.get(User, user_id)
    if not target: raise HTTPException(404, "用户不存在")
    target.role = body.role; db.commit()
    return {"id": target.id, "role": target.role}


@app.put("/api/admin/users/{user_id}/courses")
def set_user_courses(user_id: int, body: CourseAccessBody, user: User = Depends(admin_only), db: Session = Depends(get_db)):
    target = db.get(User, user_id)
    if not target: raise HTTPException(404, "用户不存在")
    if target.role == "admin": raise HTTPException(400, "管理员账户不使用学生选课权限")
    ids = list(dict.fromkeys(body.course_ids))
    existing = set(db.scalars(select(Course.id).where(Course.id.in_(ids or [-1]))).all())
    if existing != set(ids): raise HTTPException(400, "包含无效课程")
    db.execute(delete(Enrollment).where(Enrollment.user_id == user_id))
    db.add_all([Enrollment(user_id=user_id, course_id=course_id) for course_id in ids])
    db.commit()
    return {"user_id": user_id, "course_ids": ids}


def event(name: str, data: Any) -> str:
    return f"event: {name}\ndata: {json.dumps(data, ensure_ascii=False, default=str)}\n\n"


async def run_agent(message: str, user: User, db: Session) -> AsyncIterator[str]:
    yield event("status", {"stage": "thinking", "message": "正在理解你的问题"})
    llm_key = os.getenv("OPENAI_API_KEY")
    if not llm_key:
        async for item in offline_agent(message, user, db): yield item
        return
    try:
        from langchain.agents import create_agent
        from langchain_core.messages import HumanMessage
        from langchain_core.tools import tool
        from langchain_openai import ChatOpenAI

        @tool
        async def get_my_schedule() -> str:
            """查询当前登录学生本人的课程表。不得查询其他用户。"""
            result = await call_mcp_tool("get_my_schedule", {"auth_context": signed_user_context(user.id)})
            if result is not None: return result
            courses = db.scalars(select(Course).join(Enrollment).where(Enrollment.user_id == user.id).order_by(Course.weekday, Course.start_time)).all()
            return json.dumps([course_json(c) for c in courses], ensure_ascii=False)

        @tool
        async def search_course_materials(query: str) -> str:
            """在当前学生有权限访问的课程资料中进行 RAG 检索。"""
            result = await call_mcp_tool("search_course_materials", {"auth_context": signed_user_context(user.id), "query": query})
            return result if result is not None else json.dumps(await search_materials(query, user, db), ensure_ascii=False)

        @tool
        async def save_study_plan(title: str, content: str, due_date: str = "") -> str:
            """为当前登录学生保存学习计划。"""
            result = await call_mcp_tool("save_study_plan", {"auth_context": signed_user_context(user.id), "title": title, "content": content, "due_date": due_date})
            if result is not None: return result
            parsed = date.fromisoformat(due_date) if due_date else None
            row = Plan(user_id=user.id, title=title, content=content, due_date=parsed)
            db.add(row); db.commit(); db.refresh(row)
            return json.dumps({"id": row.id, "title": title, "saved": True}, ensure_ascii=False)

        tools = [get_my_schedule, search_course_materials, save_study_plan]
        llm = ChatOpenAI(model=os.getenv("OPENAI_MODEL", "deepseek-chat"), api_key=llm_key, base_url=os.getenv("OPENAI_BASE_URL", "https://api.deepseek.com/v1"), temperature=0.2, streaming=True)
        agent = create_agent(model=llm, tools=tools, system_prompt=f"你是校园课程智能助手。只回答当前用户有权限访问的信息。用户称呼：{user.display_name}。需要查课表时调用 get_my_schedule；依据课程资料回答时调用 search_course_materials；用户要求制定并保存计划时先生成计划并调用 save_study_plan。没有资料依据时明确说明。使用简体中文。")
        answer = ""
        labels = {"get_my_schedule": "查询你的课表", "search_course_materials": "检索课程资料", "save_study_plan": "保存学习计划"}
        async for item in agent.astream_events({"messages": [HumanMessage(content=message)]}, config={"recursion_limit": 12}, version="v2"):
            kind = item.get("event")
            name = item.get("name", "")
            data = item.get("data", {})
            if kind == "on_tool_start":
                yield event("tool_start", {"name": name, "label": labels.get(name, name), "input": data.get("input", {})})
            elif kind == "on_tool_end":
                result = data.get("output", "")
                if hasattr(result, "content"): result = result.content
                yield event("tool_result", {"name": name, "result": result})
            elif kind == "on_tool_error":
                yield event("tool_result", {"name": name, "result": f"工具执行失败：{str(data.get('error', ''))[:250]}"})
            elif kind == "on_chat_model_stream":
                chunk = data.get("chunk")
                text = getattr(chunk, "content", "")
                if isinstance(text, list): text = "".join(block.get("text", "") for block in text if isinstance(block, dict))
                if isinstance(text, str) and text:
                    answer += text
                    yield event("token", {"text": text})
            elif kind == "on_chain_end" and isinstance(data.get("output"), dict):
                final_messages = data["output"].get("messages", [])
                if final_messages and not answer:
                    final_content = getattr(final_messages[-1], "content", "")
                    if isinstance(final_content, str): answer = final_content
        yield event("done", {"answer": answer})
    except Exception as exc:
        yield event("error", {"message": f"AI 服务暂不可用：{str(exc)[:300]}"})


async def offline_agent(message: str, user: User, db: Session) -> AsyncIterator[str]:
    q = message.lower()
    if any(word in q for word in ("课表", "课程安排", "今天有什么课", "明天有什么课")):
        label, name = "查询你的课表", "get_my_schedule"
        yield event("tool_start", {"name": name, "label": label, "input": {}})
        mcp_result = await call_mcp_tool(name, {"auth_context": signed_user_context(user.id)})
        try: result = json.loads(mcp_result) if mcp_result is not None else None
        except Exception: result = None
        if result is None:
            rows = db.scalars(select(Course).join(Enrollment).where(Enrollment.user_id == user.id).order_by(Course.weekday, Course.start_time)).all()
            result = [course_json(c) for c in rows]
        yield event("tool_result", {"name": name, "result": result})
        days = ["", "周一", "周二", "周三", "周四", "周五", "周六", "周日"]
        answer = "这是你本人的课表：\n" + "\n".join(f"{days[c['weekday']]} {c['start_time']}-{c['end_time']}｜{c['name']}｜{c['room']}" for c in result)
    elif any(word in q for word in ("计划", "规划", "安排学习", "复习计划")):
        title = "本周学习计划"
        tasks = db.scalars(select(Task).where(Task.user_id == user.id, Task.completed.is_(False)).order_by(Task.due_date)).all()
        lines = [f"- 优先完成「{t.title}」（{t.due_date.isoformat()} 截止）" for t in tasks[:3]] or ["- 整理本周课程笔记并复习重点"]
        content = "\n".join(lines + ["- 每天预留 30 分钟回顾课堂内容", "- 完成后检查并记录疑难问题"])
        yield event("tool_start", {"name": "save_study_plan", "label": "保存学习计划", "input": {"title": title}})
        due = date.today() + timedelta(days=7)
        mcp_result = await call_mcp_tool("save_study_plan", {"auth_context": signed_user_context(user.id), "title": title, "content": content, "due_date": due.isoformat()})
        if mcp_result is None:
            row = Plan(user_id=user.id, title=title, content=content, due_date=due); db.add(row); db.commit(); db.refresh(row)
            saved = {"id": row.id, "saved": True}
        else:
            try: saved = json.loads(mcp_result)
            except Exception: saved = {"saved": True}
        yield event("tool_result", {"name": "save_study_plan", "result": saved})
        answer = f"已为你生成并保存「{title}」：\n{content}"
    else:
        yield event("tool_start", {"name": "search_course_materials", "label": "检索课程资料", "input": {"query": message}})
        mcp_result = await call_mcp_tool("search_course_materials", {"auth_context": signed_user_context(user.id), "query": message})
        try: rows = json.loads(mcp_result) if mcp_result is not None else None
        except Exception: rows = None
        if rows is None: rows = await search_materials(message, user, db)
        yield event("tool_result", {"name": "search_course_materials", "result": rows})
        if rows:
            answer = "根据你的课程资料：\n\n" + "\n\n".join(f"《{r['filename']}》\n{r['content']}" for r in rows[:3])
        else:
            answer = "我没有在你有权限访问的课程资料中找到相关内容。可以请管理员上传对应课程资料后再问我。"
    for i in range(0, len(answer), 28):
        await asyncio.sleep(0.025)
        yield event("token", {"text": answer[i:i+28]})
    yield event("done", {"answer": answer})


@app.post("/api/chat/stream")
async def chat_stream(body: ChatBody, user: User = Depends(get_user)):
    async def stream():
        with SessionLocal() as db:
            async for item in run_agent(body.message, user, db): yield item
    return StreamingResponse(stream(), media_type="text/event-stream", headers={"Cache-Control": "no-cache", "X-Accel-Buffering": "no"})


@app.on_event("startup")
async def index_seed_docs():
    if os.getenv("EMBEDDING_API_KEY"):
        with SessionLocal() as db:
            for doc in db.scalars(select(Document)).all(): await vector_add(doc)

