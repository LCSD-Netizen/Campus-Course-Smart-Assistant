"""课程资料解析、切分、向量化与检索服务。"""

from __future__ import annotations
import asyncio
import hashlib
import hmac
import json
import os
import re
import secrets
import uuid
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
from typing import Any, AsyncIterator
import httpx
import jwt
from dotenv import load_dotenv
from fastapi import Depends, FastAPI, File, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse
from fastapi.security import HTTPAuthorizationCredentials, HTTPBearer
from pydantic import BaseModel, Field
from sqlalchemy import Boolean, Date, DateTime, ForeignKey, Integer, String, Text, create_engine, delete, func, select
from sqlalchemy.orm import DeclarativeBase, Mapped, Session, mapped_column, relationship, sessionmaker
from backend.core.shared import *

def split_chunks(text: str, size: int=450, overlap: int=80) -> list[str]:
    """处理split_chunks相关业务逻辑。"""
    text = re.sub('\\s+', ' ', text).strip()
    return [text[i:i + size] for i in range(0, len(text), size - overlap) if text[i:i + size].strip()]


async def embeddings(texts: list[str]) -> list[list[float]] | None:
    """处理embeddings相关业务逻辑。"""
    key = os.getenv('EMBEDDING_API_KEY')
    if not key:
        return None
    url = os.getenv('EMBEDDING_BASE_URL', 'https://api.siliconflow.cn/v1').rstrip('/') + '/embeddings'
    try:
        async with httpx.AsyncClient(timeout=45) as client:
            response = await client.post(url, headers={'Authorization': f'Bearer {key}'}, json={'model': os.getenv('EMBEDDING_MODEL', 'BAAI/bge-m3'), 'input': texts})
            response.raise_for_status()
            return [item['embedding'] for item in response.json()['data']]
    except Exception:
        return None


async def vector_add(doc: Document):
    """处理vector_add相关业务逻辑。"""
    chunks = split_chunks(doc.content)
    vectors = await embeddings(chunks)
    if vectors is None:
        return
    try:
        import chromadb
        path = Path(os.getenv('CHROMA_DIR', str(DATA_DIR / 'chroma')))
        path.mkdir(parents=True, exist_ok=True)
        client = chromadb.PersistentClient(path=str(path))
        collection = client.get_or_create_collection('campus_course_materials')
        collection.add(ids=[f'doc-{doc.id}-{i}' for i in range(len(chunks))], documents=chunks, embeddings=vectors, metadatas=[{'document_id': doc.id, 'course_id': doc.course_id or 0, 'filename': doc.filename} for _ in chunks])
    except Exception:
        return


async def search_materials(query: str, user: User, db: Session, limit: int=4) -> list[dict[str, Any]]:
    """处理search_materials相关业务逻辑。"""
    enrolled = db.scalars(select(Enrollment.course_id).where(Enrollment.user_id == user.id)).all()
    docs = db.scalars(select(Document).where(Document.course_id.is_(None) | Document.course_id.in_(enrolled if enrolled else [-1]))).all()
    query_vector = await embeddings([query])
    if query_vector:
        try:
            import chromadb
            client = chromadb.PersistentClient(path=os.getenv('CHROMA_DIR', str(DATA_DIR / 'chroma')))
            collection = client.get_or_create_collection('campus_course_materials')
            result = collection.query(query_embeddings=query_vector, n_results=limit, where={'course_id': {'$in': enrolled}} if enrolled else {'course_id': -1})
            doc_map = {d.id: d for d in docs}
            ids = result.get('ids', [[]])[0]
            chunks = result.get('documents', [[]])[0]
            rows = []
            for chunk_id, chunk in zip(ids, chunks):
                match = re.search('doc-(\\d+)-', chunk_id)
                document_id = int(match.group(1)) if match else None
                if document_id in doc_map:
                    rows.append({'filename': doc_map[document_id].filename, 'content': chunk})
            if rows:
                return rows[:limit]
        except Exception:
            pass
    terms = [t.lower() for t in re.findall('[\\w\\u4e00-\\u9fff]+', query) if len(t) > 1]
    ranked = sorted(docs, key=lambda d: sum((d.content.lower().count(t) + d.filename.lower().count(t) for t in terms)), reverse=True)
    return [{'filename': d.filename, 'content': d.content[:650]} for d in ranked[:limit]]


async def extract_text(filename: str, data: bytes) -> str:
    """处理extract_text相关业务逻辑。"""
    suffix = Path(filename).suffix.lower()
    if suffix == '.pdf':
        from pypdf import PdfReader
        import io
        return '\n'.join((page.extract_text() or '' for page in PdfReader(io.BytesIO(data)).pages))
    if suffix in ('.docx',):
        from docx import Document as WordDocument
        import io
        return '\n'.join((p.text for p in WordDocument(io.BytesIO(data)).paragraphs))
    if suffix in ('.pptx',):
        from pptx import Presentation
        import io
        prs = Presentation(io.BytesIO(data))
        return '\n'.join((shape.text for slide in prs.slides for shape in slide.shapes if getattr(shape, 'has_text_frame', False)))
    if suffix == '.txt':
        return data.decode('utf-8', errors='ignore')
    raise HTTPException(400, '支持 PDF、DOCX、PPTX、TXT 格式')